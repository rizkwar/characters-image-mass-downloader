"""Download image search results for character names in a CSV or XLSX file."""

from __future__ import annotations

import argparse
import csv
import re
import sys
import time
from pathlib import Path
from urllib.parse import urlparse

import requests
from ddgs import DDGS
from openpyxl import load_workbook


IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".gif", ".bmp"}


def read_names(input_path: Path, column: str) -> list[str]:
    """Read unique, non-empty character names from CSV or XLSX."""
    suffix = input_path.suffix.lower()
    if suffix == ".csv":
        with input_path.open("r", encoding="utf-8-sig", newline="") as file:
            rows = csv.DictReader(file)
            if not rows.fieldnames:
                raise ValueError("The CSV must have a header row.")
            field = column if column in rows.fieldnames else rows.fieldnames[0]
            values = [row.get(field, "") for row in rows]
    elif suffix in {".xlsx", ".xlsm"}:
        workbook = load_workbook(input_path, read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(rows, ())]
        if not headers:
            raise ValueError("The spreadsheet must have a header row.")
        index = headers.index(column) if column in headers else 0
        values = [row[index] if len(row) > index else "" for row in rows]
        workbook.close()
    else:
        raise ValueError("Input must be a .csv, .xlsx, or .xlsm file.")

    seen: set[str] = set()
    names: list[str] = []
    for value in values:
        name = str(value).strip() if value is not None else ""
        if name and name.casefold() not in seen:
            names.append(name)
            seen.add(name.casefold())
    return names


def safe_folder_name(name: str) -> str:
    """Make a character name safe to use as a Windows folder name."""
    cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", name).strip(" .")
    return cleaned or "unnamed"


def extension_for(url: str, content_type: str) -> str:
    suffix = Path(urlparse(url).path).suffix.lower()
    if suffix in IMAGE_EXTENSIONS:
        return ".jpg" if suffix == ".jpeg" else suffix
    content_suffix = content_type.split("/")[-1].split(";")[0].lower()
    return {"jpeg": ".jpg", "jpg": ".jpg", "png": ".png", "webp": ".webp", "gif": ".gif", "bmp": ".bmp"}.get(content_suffix, ".jpg")


def download_character_images(name: str, output_dir: Path, count: int, safe_search: bool, delay: float) -> tuple[int, str | None]:
    folder = output_dir / safe_folder_name(name)
    folder.mkdir(parents=True, exist_ok=True)
    saved = 0
    try:
        query = f"{name} character"
        with DDGS() as search:
            results = list(search.images(query, max_results=count * 3, safesearch="on" if safe_search else "off"))
        session = requests.Session()
        session.headers.update({"User-Agent": "image-mass-downloader/1.0"})
        for result in results:
            if saved >= count:
                break
            image_url = result.get("image")
            if not image_url:
                continue
            try:
                response = session.get(image_url, timeout=20, stream=True)
                response.raise_for_status()
                content_type = response.headers.get("Content-Type", "")
                if not content_type.startswith("image/"):
                    continue
                extension = extension_for(image_url, content_type)
                target = folder / f"{saved + 1:02d}{extension}"
                with target.open("wb") as file:
                    for chunk in response.iter_content(chunk_size=65536):
                        if chunk:
                            file.write(chunk)
                saved += 1
                time.sleep(delay)
            except requests.RequestException:
                continue
        return saved, None
    except Exception as error:  # Keep one failed search from stopping the batch.
        return saved, str(error)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Download image search results for names in a spreadsheet.")
    parser.add_argument("input", type=Path, help="CSV or XLSX file containing names")
    parser.add_argument("--column", default="name", help="Column header to read (default: name)")
    parser.add_argument("--output", type=Path, default=Path("downloads"), help="Output folder (default: downloads)")
    parser.add_argument("--count", type=int, default=5, help="Images per name (default: 5)")
    parser.add_argument("--delay", type=float, default=0.5, help="Delay between downloads in seconds (default: 0.5)")
    parser.add_argument("--unsafe", action="store_true", help="Disable safe-search filtering")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    if args.count < 1:
        print("--count must be at least 1.", file=sys.stderr)
        return 2
    try:
        names = read_names(args.input, args.column)
    except (OSError, ValueError) as error:
        print(f"Could not read input: {error}", file=sys.stderr)
        return 1
    if not names:
        print("No names found in the input file.", file=sys.stderr)
        return 1

    args.output.mkdir(parents=True, exist_ok=True)
    report = args.output / "download_report.csv"
    print(f"Downloading {args.count} image(s) for {len(names)} character(s)...")
    with report.open("w", encoding="utf-8", newline="") as file:
        writer = csv.writer(file)
        writer.writerow(["name", "downloaded", "error"])
        for number, name in enumerate(names, start=1):
            saved, error = download_character_images(name, args.output, args.count, not args.unsafe, args.delay)
            writer.writerow([name, saved, error or ""])
            print(f"[{number}/{len(names)}] {name}: {saved} image(s)" + (f" - {error}" if error else ""))

    print(f"Finished. Files are in {args.output.resolve()}")
    print(f"Report: {report.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
